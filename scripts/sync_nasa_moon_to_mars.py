#!/usr/bin/env python3
"""Build the Moon→Mars Decision Atlas dataset from NASA's official XLSX products.

Discovers the current 2025 Architecture Concept Review spreadsheets from NASA's
Architecture Definition Documents page, downloads them, records SHA-256 provenance,
normalizes workbook rows, validates the result, and atomically writes JSON.

Fail-closed behavior: if discovery, download, parsing, or validation fails, the
existing dataset remains untouched.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup, Tag
from openpyxl import load_workbook

ARCHITECTURE_PAGE = "https://www.nasa.gov/moontomarsarchitecture-architecturedefinitiondocuments/"
NTRS_SOURCE = "https://ntrs.nasa.gov/citations/20250010956"
DOCUMENT_ID = "20250010956"
REVISION = "C"
ACR = "2025 ACR"
EXPECTED_KINDS = (
    "lunar_objectives",
    "mars_objectives",
    "technology_gaps",
    "data_gaps",
)

SEGMENTS = [
    "Human Lunar Return",
    "Foundational Exploration",
    "Sustained Lunar Evolution",
    "Humans to Mars",
]

SUBARCHITECTURES = [
    "Autonomous Systems & Robotics",
    "Communications, Positioning, Navigation & Timing",
    "Data Systems and Management",
    "Habitation Systems",
    "Human Systems",
    "In-situ Resource Utilization Systems",
    "Infrastructure Support",
    "Logistics Systems",
    "Mobility Systems",
    "Power Systems",
    "Transportation Systems",
    "Utilization Systems",
]

KNOWN_TOP_IDS = {"0801", "0201", "0301", "1107", "0103", "1104"}


@dataclass(frozen=True)
class SourceFile:
    kind: str
    url: str
    filename: str
    sha256: str
    content: bytes


def get(url: str, *, timeout: int = 45) -> requests.Response:
    response = requests.get(
        url,
        timeout=timeout,
        headers={
            "User-Agent": "MoonMarsDecisionAtlas/1.0 (+https://github.com/oluwafemidiakhoa/nasa-spa)"
        },
    )
    response.raise_for_status()
    return response


def is_xlsx_href(href: str | None) -> bool:
    return bool(href and re.search(r"\.xlsx(?:$|[?#])", href, re.I))


def heading_text(tag: Tag) -> str:
    return " ".join(tag.stripped_strings).strip()


def find_2025_section(soup: BeautifulSoup) -> list[Tag]:
    start = None
    for tag in soup.find_all(re.compile(r"^h[1-6]$")):
        if "2025 Architecture Concept Review" in heading_text(tag):
            start = tag
            break
    if start is None:
        raise RuntimeError("Could not locate the 2025 Architecture Concept Review section")

    tags: list[Tag] = []
    for node in start.next_elements:
        if not isinstance(node, Tag):
            continue
        if node.name and re.match(r"^h[1-6]$", node.name):
            text = heading_text(node)
            if node is not start and "2024 Architecture Concept Review" in text:
                break
        tags.append(node)
    return tags


def nearby_text(anchor: Tag, *, chars: int = 1200) -> str:
    chunks: list[str] = []
    parent = anchor
    for _ in range(5):
        parent = parent.parent if isinstance(parent.parent, Tag) else None
        if parent is None:
            break
        text = " ".join(parent.stripped_strings)
        if text:
            chunks.append(text)
        if sum(map(len, chunks)) >= chars:
            break
    return " ".join(chunks)[:chars].lower()


def classify_link(url: str, context: str) -> str | None:
    hay = (url + " " + context).lower()
    if "lunar" in hay and ("objective" in hay or "mapping" in hay):
        return "lunar_objectives"
    if "mars" in hay and ("objective" in hay or "mapping" in hay):
        return "mars_objectives"
    if "technology" in hay and "gap" in hay:
        return "technology_gaps"
    if "data" in hay and "gap" in hay:
        return "data_gaps"
    return None


def discover_xlsx_urls(page_url: str = ARCHITECTURE_PAGE) -> dict[str, str]:
    soup = BeautifulSoup(get(page_url).text, "html.parser")
    section = find_2025_section(soup)
    anchors: list[Tag] = []
    seen: set[str] = set()

    for tag in section:
        if tag.name != "a":
            continue
        href = tag.get("href")
        if not is_xlsx_href(href):
            continue
        url = urljoin(page_url, href)
        if url not in seen:
            seen.add(url)
            anchors.append(tag)

    # Defensive fallback if NASA changes the card markup.
    if len(anchors) < 4:
        anchors = []
        seen.clear()
        for tag in soup.find_all("a", href=True):
            href = tag.get("href")
            if not is_xlsx_href(href):
                continue
            url = urljoin(page_url, href)
            if url not in seen:
                seen.add(url)
                anchors.append(tag)

    classified: dict[str, str] = {}
    for tag in anchors:
        url = urljoin(page_url, tag.get("href"))
        kind = classify_link(url, nearby_text(tag))
        if kind and kind not in classified:
            classified[kind] = url

    # NASA currently presents the four 2025 XLSX products in this order. Use that
    # ordering only if card context/filename classification was insufficient.
    unique_urls = list(dict.fromkeys(urljoin(page_url, a.get("href")) for a in anchors))
    if len(unique_urls) >= 4:
        for kind, url in zip(EXPECTED_KINDS, unique_urls[:4]):
            classified.setdefault(kind, url)

    missing = [kind for kind in EXPECTED_KINDS if kind not in classified]
    if missing:
        raise RuntimeError(
            "Could not discover all 2025 NASA XLSX products. "
            f"Missing: {', '.join(missing)}; found {len(unique_urls)} XLSX links"
        )
    return {kind: classified[kind] for kind in EXPECTED_KINDS}


def download_sources(urls: dict[str, str]) -> dict[str, SourceFile]:
    out: dict[str, SourceFile] = {}
    for kind, url in urls.items():
        content = get(url).content
        if len(content) < 1000:
            raise RuntimeError(f"Downloaded {kind} file is unexpectedly small: {len(content)} bytes")
        out[kind] = SourceFile(
            kind=kind,
            url=url,
            filename=Path(urlparse(url).path).name or f"{kind}.xlsx",
            sha256=hashlib.sha256(content).hexdigest(),
            content=content,
        )
    return out


def scalar(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (int, float, bool)):
        return value
    return str(value).strip()


def header_key(value: Any) -> str:
    text = str(value or "").strip().lower().replace("&", " and ")
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def find_header_row(rows: list[tuple[Any, ...]]) -> int:
    keywords = {
        "id", "gap", "title", "objective", "description", "priority", "segment",
        "sub architecture", "data", "utility", "function", "use case", "traceability",
    }
    best_idx, best_score = 0, -1.0
    for idx, row in enumerate(rows[:35]):
        vals = [header_key(v) for v in row if v not in (None, "")]
        if len(vals) < 2:
            continue
        keyword_hits = sum(any(keyword in value for keyword in keywords) for value in vals)
        score = len(vals) + keyword_hits * 3
        if score > best_score:
            best_idx, best_score = idx, score
    return best_idx


def unique_headers(row: tuple[Any, ...]) -> list[str]:
    headers: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(row, start=1):
        base = str(value).strip() if value not in (None, "") else f"column_{index}"
        counts[base] = counts.get(base, 0) + 1
        headers.append(base if counts[base] == 1 else f"{base}__{counts[base]}")
    return headers


def workbook_records(source: SourceFile) -> list[dict[str, Any]]:
    workbook = load_workbook(io.BytesIO(source.content), read_only=True, data_only=True)
    records: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        header_idx = find_header_row(rows)
        headers = unique_headers(rows[header_idx])
        blank_streak = 0
        for excel_row, values in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
            vals = [scalar(value) for value in values]
            if not any(value not in (None, "") for value in vals):
                blank_streak += 1
                if blank_streak >= 25:
                    break
                continue
            blank_streak = 0
            record = {
                headers[i]: vals[i] if i < len(vals) else None
                for i in range(len(headers))
            }
            record["_source_sheet"] = sheet.title
            record["_source_row"] = excel_row
            records.append(record)
    return records


def normalized_items(record: dict[str, Any]) -> dict[str, Any]:
    return {header_key(key): value for key, value in record.items() if not key.startswith("_")}


def pick(record: dict[str, Any], exact: Iterable[str] = (), contains: Iterable[str] = ()) -> Any:
    norm = normalized_items(record)
    for candidate in exact:
        if candidate in norm and norm[candidate] not in (None, ""):
            return norm[candidate]
    for needle in contains:
        for key, value in norm.items():
            if needle in key and value not in (None, ""):
                return value
    return None


def as_int(value: Any) -> int | None:
    if value in (None, ""):
        return None
    match = re.search(r"-?\d+", str(value))
    return int(match.group()) if match else None


def gap_id(value: Any) -> str | None:
    if value in (None, ""):
        return None
    match = re.search(r"(?<!\d)(\d{4})(?!\d)", str(value))
    return match.group(1) if match else None


def split_list(value: Any) -> list[str]:
    if value in (None, ""):
        return []
    return [part.strip() for part in re.split(r"[;,|\n]+", str(value)) if part.strip()]


def truthy_marker(value: Any) -> bool:
    if value in (None, "", 0, False):
        return False
    return str(value).strip().lower() not in {"no", "n", "false", "0", "none", "n/a", "na", "-"}


def taxonomy_hits(
    record: dict[str, Any], official_names: list[str], aggregate_needles: list[str]
) -> list[str]:
    norm = normalized_items(record)
    hits: list[str] = []
    for official in official_names:
        official_key = header_key(official)
        for key, value in norm.items():
            if key == official_key and truthy_marker(value):
                hits.append(official)
                break
    for needle in aggregate_needles:
        value = pick(record, exact=[needle], contains=[needle])
        for token in split_list(value):
            token_key = header_key(token)
            for official in official_names:
                official_key = header_key(official)
                if token_key == official_key or token_key in official_key or official_key in token_key:
                    if official not in hits:
                        hits.append(official)
    return hits


def normalize_technology_gaps(
    records: list[dict[str, Any]], source: SourceFile
) -> list[dict[str, Any]]:
    gaps: list[dict[str, Any]] = []
    seen: set[str] = set()
    for record in records:
        gid = gap_id(pick(record, exact=["gap id", "technology gap id", "id"], contains=["gap id"]))
        title = pick(record, exact=["gap title", "technology gap title", "title"], contains=["gap title"])
        if gid is None:
            combined = pick(record, contains=["technology gap", "gap"])
            gid = gap_id(combined)
            if title in (None, "") and combined:
                title = re.sub(r"^\s*\d{4}\s*[-–—:]?\s*", "", str(combined)).strip()
        if gid is None or title in (None, "") or gid in seen:
            continue
        seen.add(gid)

        rank = as_int(
            pick(
                record,
                exact=[
                    "priority rating",
                    "priority ranking",
                    "priority rank",
                    "overall prioritization rating",
                ],
                contains=["priority rating", "priority ranking"],
            )
        )
        priority_bin = as_int(pick(record, exact=["priority bin"], contains=["priority bin"]))
        segments = taxonomy_hits(record, SEGMENTS, ["segments", "campaign segments", "segment"])
        subarchitectures = taxonomy_hits(
            record,
            SUBARCHITECTURES,
            ["sub architectures", "subarchitecture", "sub architectures mapped"],
        )
        description = pick(
            record,
            exact=["gap description", "description"],
            contains=["gap description", "description"],
        )
        impact = pick(
            record,
            exact=["architecture impact", "impact if unresolved", "impact benefit"],
            contains=["architecture impact", "impact if", "consequence"],
        )
        current_state = pick(
            record,
            exact=["current state of the art", "current state", "state of the art"],
            contains=["current state", "state of the art"],
        )
        target = pick(
            record,
            exact=["performance target", "target performance", "desired performance"],
            contains=["performance target", "target performance"],
        )
        gaps.append(
            {
                "id": gid,
                "title": str(title).strip(),
                "priority_rank": rank,
                "priority_bin": priority_bin,
                "detail_status": (
                    "verified"
                    if any([description, impact, current_state, target, segments, subarchitectures])
                    else "rank_verified"
                ),
                "description": scalar(description),
                "impact_if_unresolved": scalar(impact),
                "current_state": scalar(current_state),
                "performance_target": scalar(target),
                "segments": segments,
                "subarchitectures": subarchitectures,
                "source_url": source.url,
                "source_workbook": source.filename,
                "source_sha256": source.sha256,
                "source_sheet": record.get("_source_sheet"),
                "source_row": record.get("_source_row"),
            }
        )
    gaps.sort(
        key=lambda gap: (
            gap["priority_rank"] is None,
            gap["priority_rank"] or 9999,
            gap["id"],
        )
    )
    return gaps


def normalize_data_gaps(
    records: list[dict[str, Any]], source: SourceFile
) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    seen: set[str] = set()
    for record in records:
        raw_id = pick(record, exact=["id", "data gap id", "gap id"], contains=["data gap id"])
        match = re.search(r"\b([A-Z]{1,3}-?\d{2,4})\b", str(raw_id or ""), re.I)
        if not match:
            continue
        gid = match.group(1).upper()
        if gid in seen:
            continue
        seen.add(gid)
        title = pick(record, exact=["data gap", "gap title", "title"], contains=["data gap"])
        utility = pick(record, exact=["data utility", "utility"], contains=["data utility"])
        out.append(
            {
                "id": gid,
                "title": scalar(title),
                "utility": scalar(utility),
                "source_url": source.url,
                "source_workbook": source.filename,
                "source_sha256": source.sha256,
                "source_sheet": record.get("_source_sheet"),
                "source_row": record.get("_source_row"),
            }
        )
    return out


def normalize_objectives(
    records: list[dict[str, Any]], source: SourceFile, domain: str
) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for record in records:
        norm = normalized_items(record)
        signal = sum(
            any(keyword in key for keyword in ("objective", "use case", "function", "goal"))
            for key in norm
        )
        if signal == 0:
            continue
        fields = {
            key: scalar(value)
            for key, value in record.items()
            if not key.startswith("_") and value not in (None, "")
        }
        if not fields:
            continue
        out.append(
            {
                "domain": domain,
                "fields": fields,
                "source_url": source.url,
                "source_workbook": source.filename,
                "source_sha256": source.sha256,
                "source_sheet": record.get("_source_sheet"),
                "source_row": record.get("_source_row"),
            }
        )
    return out


def build_dataset(
    sources: dict[str, SourceFile]
) -> tuple[dict[str, Any], dict[str, Any]]:
    tables = {kind: workbook_records(source) for kind, source in sources.items()}
    gaps = normalize_technology_gaps(tables["technology_gaps"], sources["technology_gaps"])
    data_gaps = normalize_data_gaps(tables["data_gaps"], sources["data_gaps"])
    lunar = normalize_objectives(
        tables["lunar_objectives"], sources["lunar_objectives"], "lunar"
    )
    mars = normalize_objectives(
        tables["mars_objectives"], sources["mars_objectives"], "mars"
    )

    ids = {gap["id"] for gap in gaps}
    if len(gaps) < 16 or len(ids & KNOWN_TOP_IDS) < 5:
        raise RuntimeError(
            f"Technology-gap normalization failed validation: {len(gaps)} gaps; "
            f"recognized top IDs={sorted(ids & KNOWN_TOP_IDS)}"
        )
    if not data_gaps:
        raise RuntimeError("Data-gap workbook produced zero normalized rows")
    if len(lunar) < 5 or len(mars) < 5:
        raise RuntimeError(
            f"Objective mapping parse looks incomplete: lunar={len(lunar)}, mars={len(mars)}"
        )

    generated_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    dataset = {
        "dataset": {
            "title": "NASA Moon to Mars Architecture — official spreadsheet sync",
            "document_id": DOCUMENT_ID,
            "revision": REVISION,
            "architecture_concept_review": ACR,
            "generated_at": generated_at,
            "source_url": NTRS_SOURCE,
            "architecture_page": ARCHITECTURE_PAGE,
            "generated_from_official_xlsx": True,
            "note": (
                "Generated deterministically from NASA's official 2025 Architecture "
                "Concept Review XLSX products. No missing relationship is inferred."
            ),
        },
        "segments": SEGMENTS,
        "subarchitectures": SUBARCHITECTURES,
        "technology_gaps": gaps,
        "data_gaps": data_gaps,
        "objective_mappings": {"lunar": lunar, "mars": mars},
    }
    manifest = {
        "generated_at": generated_at,
        "architecture_page": ARCHITECTURE_PAGE,
        "document_id": DOCUMENT_ID,
        "revision": REVISION,
        "sources": {
            kind: {
                "url": source.url,
                "filename": source.filename,
                "sha256": source.sha256,
                "bytes": len(source.content),
                "record_count": len(tables[kind]),
            }
            for kind, source in sources.items()
        },
        "normalized_counts": {
            "technology_gaps": len(gaps),
            "data_gaps": len(data_gaps),
            "lunar_objective_rows": len(lunar),
            "mars_objective_rows": len(mars),
        },
    }
    return dataset, manifest


def atomic_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(
        json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    temporary.replace(path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", default="data/moon_to_mars_revision_c.json")
    parser.add_argument("--manifest", default="data/moon_to_mars_source_manifest.json")
    parser.add_argument("--architecture-page", default=ARCHITECTURE_PAGE)
    args = parser.parse_args()

    try:
        urls = discover_xlsx_urls(args.architecture_page)
        print("Discovered official NASA XLSX products:")
        for kind, url in urls.items():
            print(f"  {kind}: {url}")
        sources = download_sources(urls)
        dataset, manifest = build_dataset(sources)
        atomic_json(Path(args.output), dataset)
        atomic_json(Path(args.manifest), manifest)
        print("Validation passed.")
        print(json.dumps(manifest["normalized_counts"], indent=2))
        return 0
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
